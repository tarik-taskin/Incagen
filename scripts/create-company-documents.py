"""Bundled artifact runtime ile kurumsal v2 örneğinin PDF/DOCX kaynaklarını üretir; uygulama bağımlılığı değildir."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
import csv, json, hashlib, io, zipfile, argparse
from datetime import datetime, timezone
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'fixtures/company-example/v2'
D = Decimal
sales = list(csv.DictReader((ROOT/'atlas-satis-2025.csv').open()))
budgets = list(csv.DictReader((ROOT/'atlas-butce-2025.csv').open()))
ops = list(csv.DictReader((ROOT/'atlas-operasyon-2025.csv').open()))
months = ['Ocak','Şubat','Mart','Nisan','Mayıs','Haziran','Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık']

# Golden cevapları uygulamanın TypeScript hesaplarından bağımsız Decimal aritmetiğiyle dondurur.
def totals(lo=1, hi=12):
    s=[r for r in sales if lo<=int(r['period'][-2:])<=hi]
    b=[r for r in budgets if lo<=int(r['period'][-2:])<=hi]
    o=[r for r in ops if lo<=int(r['period'][-2:])<=hi]
    sums=lambda key: sum((D(r[key]) for r in s), D(0))
    net=sums('gross_sales_try')-sums('discount_try')-sums('refund_try')
    profit=net-sums('cost_of_goods_try')
    result=dict(net=float(net),gross=float(sums('gross_sales_try')),discounts=float(sums('discount_try')),refunds=float(sums('refund_try')),profit=float(profit),budget=float(sum(D(r['net_sales_budget_try']) for r in b)),orders=sum(int(r['orders']) for r in s),contribution=float(profit-sums('marketing_try')-sums('fulfillment_try')),margin=float(profit/net*100),onTime=float(sum(D(r['on_time_orders']) for r in o)/sum(D(r['shipped_orders']) for r in o)*100),rows=len(s)+len(b)+len(o))
    for key in ['category','channel','region']:
        groups=defaultdict(Decimal)
        for r in s: groups[r[key]]+=D(r['gross_sales_try'])-D(r['discount_try'])-D(r['refund_try'])
        result[key]={k:float(v) for k,v in groups.items()}
    return result
G={k:totals(*v) for k,v in {'year':(1,12),'h1':(1,6),'h2':(7,12),'q4':(10,12)}.items()}
(ROOT/'expected.json').write_text(json.dumps(G,ensure_ascii=False,indent=2)+'\n')
M=[totals(i,i) for i in range(1,13)]
g=G['year']

# Belge ve tablo değerlerini Türkçe yerel biçimde, kuruş hassasiyetiyle yazar.
def num(value, digits=2):
    s=f'{value:,.{digits}f}'
    return s.replace(',','X').replace('.',',').replace('X','.')
def money(value): return num(value)+' TL'

PDF='atlas-performans-2025.pdf'
DOCX='atlas-ticari-degerlendirme-2025.docx'
parser=argparse.ArgumentParser();parser.add_argument('--font-dir', required=True);args=parser.parse_args()
fontroot=Path(args.font_dir)
pdfmetrics.registerFont(TTFont('Atlas',str(fontroot/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('AtlasBold',str(fontroot/'DejaVuSans-Bold.ttf')))
PAGE=(595.28,841.89)
c=canvas.Canvas(str(ROOT/PDF),pagesize=PAGE,invariant=1)
c.setTitle('Atlas Perakende 2025 Performans Raporu')
c.setAuthor('Atlas Perakende | Örnek şirket dosyası')
ink='#251B38'; muted='#6C647D'; purple='#60429B'; line='#E3DFEC'
body=ParagraphStyle('body',fontName='Atlas',fontSize=10,leading=17,textColor=colors.HexColor(ink))

# PDF sayfalarında hiyerarşiyi ve alt bilgi konumunu sabit tutar.
def start(index, eyebrow, title, desc):
    c.setFillColor(colors.HexColor('#F6F5FA'));c.rect(0,0,*PAGE,fill=1,stroke=0)
    c.setFillColor(colors.white);c.roundRect(26,49,543,737,16,fill=1,stroke=0)
    c.setFillColor(colors.HexColor(purple));c.setFont('AtlasBold',9);c.drawString(46,808,'ATLAS PERAKENDE')
    c.setFillColor(colors.HexColor(muted));c.setFont('Atlas',8);c.drawRightString(549,808,'2025 / Yönetim dosyası')
    c.setFont('AtlasBold',8);c.setFillColor(colors.HexColor(purple));c.drawString(46,756,eyebrow)
    c.setFont('AtlasBold',24);c.setFillColor(colors.HexColor(ink));c.drawString(46,717,title)
    p=Paragraph(escape(desc),body);w,h=p.wrap(499,80);p.drawOn(c,46,691-h)
    c.setFont('Atlas',8);c.setFillColor(colors.HexColor(muted));c.drawString(46,28,'Atlas Perakende / 2025 performans ve kaynak dosyası');c.drawRightString(549,28,f'{index:02d} / 07')

def para(text,y):
    p=Paragraph(escape(text),body);_,h=p.wrap(499,200);p.drawOn(c,46,y-h);return y-h-16

def table(headers, rows, y, widths=None):
    style=ParagraphStyle('cell',fontName='Atlas',fontSize=8,leading=12)
    data=[[Paragraph(escape(str(x)),style) for x in headers]]+[[Paragraph(escape(str(x)),style) for x in r] for r in rows]
    t=Table(data,colWidths=widths or [499/len(headers)]*len(headers),repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EDE9F5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F8F7FB')]),('GRID',(0,0),(-1,-1),0.4,colors.HexColor(line)),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    _,h=t.wrap(499,650);t.drawOn(c,46,y-h);return y-h-18

start(1,'01 / YÖNETİCİ ÖZETİ','Yıllık performans görünümü','Dönem: 1 Ocak - 31 Aralık 2025. Yönetim ekibi için satış, bütçe, brüt kârlılık ve hizmet kalitesinin ortak değerlendirmesi.')
y=table(['Net satış','Brüt kâr','Bütçe gerçekleşmesi'],[[money(g['net']),money(g['profit']),'%'+num(g['net']/g['budget']*100)]],610,[172,172,155])
y=para(f"Yıl {money(g['net'])} net satış ve {num(g['orders'],0)} sipariş ile kapanmıştır. Brüt kâr marjı %{num(g['margin'])}; pazarlama ve sipariş karşılama giderleri sonrası katkı {money(g['contribution'])} düzeyindedir. Katkı, faaliyet kârı veya net kâr olarak yorumlanmamalıdır.",y)
y=para('Ticari görünüm dört bölge, üç kanal, sekiz ürün kategorisi ve üç müşteri segmentini kapsar. Satış dosyasındaki 3.456 satır aylık kırılım toplamıdır; her satır tekil bir satış işlemi değildir. Siparişler bu kırılımlarda ayrık tanımlanmıştır.',y)
y=para('Yönetim önceliği gelir büyüklüğünü bütçe disiplini ve hizmet kapasitesiyle birlikte değerlendirmektir. Kanal yatırımı kararı için brüt marj, sipariş başına katkı ve karşılanabilen talep birlikte incelenmelidir. Bu paket kampanyaların nedensel etkisini ölçmez.',y)
y=table(['Kaynak','Kapsam'],[['Satış kayıtları','3.456 satır / 14 alan'],['Satış bütçesi','12 aylık hedef'],['Teslimat operasyonu','144 bölge ve kanal kaydı'],['Ticari değerlendirme','Kanal, portföy ve kontrol kararları']],y,[178,321])
c.showPage()
start(2,'02 / DÖNEMSEL PERFORMANS','Aylık gelir ve bütçe','Aynı ayların net satış ve bütçeleri karşılaştırılır. Birimler TL ve KDV hariçtir. Yüzdeler toplamlardan hesaplanır.')
y=table(['Ay','Net satış','Bütçe','Brüt marj'],[[months[i],num(r['net']),num(r['budget']),'%'+num(r['margin'])] for i,r in enumerate(M)],610,[72,155,155,117])
y=para(f"Yıl toplamında bütçeye göre fark {money(g['net']-g['budget'])} seviyesindedir. İlk altı ay {money(G['h1']['net'])}, son altı ay {money(G['h2']['net'])} net satış üretmiştir.",y)
c.showPage()
start(3,'03 / KANAL VE BÖLGE','Gelirin dağılımı','Kanal ve bölge toplamları aynı satış kapsamının farklı kırılımlarıdır. Bu iki tablo birbirine eklenmez.')
y=table(['Kanal','Net satış','Gelir payı'],[[k,money(v),'%'+num(v/g['net']*100)] for k,v in g['channel'].items()],610,[155,210,134])
y=table(['Bölge','Net satış','Gelir payı'],[[k,money(v),'%'+num(v/g['net']*100)] for k,v in g['region'].items()],y,[155,210,134])
y=para('Web mağazası, pazaryeri ve mağaza kanallarında sipariş karşılama maliyetleri farklıdır. Gelir payının yüksek olması tek başına yatırım önceliği oluşturmaz. İade, indirim ve maliyetler her kırılımda net gelirle aynı döneme aittir.',y)
c.showPage()
start(4,'04 / ÜRÜN PORTFÖYÜ','Kategori performansı','Sekiz kategori için satış kompozisyonu. Karşılaştırma fiyat, ürün karması ve hacmi birlikte içerir; hacim etkisini tek başına ayırmaz.')
y=table(['Kategori','Net satış','Gelir payı'],[[k,money(v),'%'+num(v/g['net']*100)] for k,v in sorted(g['category'].items(),key=lambda x:-x[1])],610,[160,205,134])
y=para('Portföy yönetiminde yüksek gelirli kategoriler ile yüksek katkı üreten kategoriler ayrı izlenmelidir. Stok devir hızı, tedarik süresi ve kayıp satış verileri bu pakete dahil değildir. Ürün veya tedarikçi seviyesinde satın alma kararı için bu alanlar eklenmelidir.',y)
c.showPage()
start(5,'05 / HİZMET KALİTESİ','Teslimat ve kapasite','Her ay, bölge ve kanal için sevk edilen siparişler zamanında ve gecikmeli olarak ayrılır. Operasyon toplamı satış siparişleriyle eşleşir.')
y=table(['Ay','Sipariş','Zamanında teslimat'],[[months[i],num(r['orders'],0),'%'+num(r['onTime'])] for i,r in enumerate(M)],610,[150,175,174])
y=para(f"Yıllık zamanında teslimat %{num(g['onTime'])} olarak ölçülmüştür. Payda toplam sevk edilen sipariş; pay toplam zamanında teslim edilen sipariştir. Aylık oranların basit ortalaması kullanılmaz.",y)
c.showPage()
start(6,'06 / MALİ TANIMLAR','Gelirden katkıya','Bütün kalemler aynı para biriminde ve aynı dönem kapsamındadır. İndirim ve iadeler satır toplam tutarıdır; ürün adediyle yeniden çarpılmaz.')
sumfield=lambda key:sum(D(r[key]) for r in sales)
y=table(['Hesap kalemi','Tutar'],[['Brüt satış',money(g['gross'])],['İndirimler',money(g['discounts'])],['İadeler',money(g['refunds'])],['Net satış',money(g['net'])],['Satılan mal maliyeti',money(sumfield('cost_of_goods_try'))],['Brüt kâr',money(g['profit'])],['Pazarlama',money(sumfield('marketing_try'))],['Sipariş karşılama',money(sumfield('fulfillment_try'))],['Katkı tutarı',money(g['contribution'])]],610,[245,254])
y=para('Net satış = brüt satış - indirim - iade. Brüt kâr = net satış - satılan mal maliyeti. Katkı = brüt kâr - pazarlama - sipariş karşılama. Genel yönetim, kira, amortisman, vergi ve finansman giderleri kapsam dışında olduğundan net kâr sonucu çıkarılamaz.',y)
c.showPage()
start(7,'07 / YÖNETİM GÜNDEMİ','Karar ve takip çerçevesi','Karar önerileri gözlemlerden ayrılır. Kaynak dosyaları aynı kapsam ve dönem üzerinden okunur; ek kanıt gerektiren varsayımlar açık tutulur.')
y=610
for title,text in [('Kanal yatırımı','Kanal performansını net satış, brüt marj ve sipariş başına katkıyla birlikte değerlendirin. Pazaryeri komisyonları veya müşteri edinme maliyeti ayrı kalem olarak doğrulanmadan pazarlama bütçesi için otomatik artış kararı vermeyin.'),('Hizmet kapasitesi','Yıl sonu hacmini aylık teslimat oranlarıyla birlikte izleyin. Kanal ve bölge bazında gecikmeli siparişlerin operasyon nedenleri araştırılmalıdır; bu dosya taşıyıcı veya depo düzeyinde neden kodu içermez.'),('Veri kontrolü','Aylık siparişler satış ve operasyon dosyaları arasında uzlaştırılmıştır. Bölge, kategori ve kanal toplamları aynı net satışa döner. Bütçe gerçekleşmesi aynı aylardaki bütçeyi payda alır; dönem değiştirildiğinde iki taraf birlikte filtrelenir.'),('Ek çalışma','Bir sonraki değerlendirmeye stok, tedarik süresi, iade nedeni ve kanal komisyon dökümleri eklenebilir. Fiyat ve hacim ayrıştırması için ürün seviyesinde adet, birim fiyat ve ürün eşleştirme tablosu gerekir.')]:
    c.setFont('AtlasBold',12);c.setFillColor(colors.HexColor(ink));c.drawString(46,y,title);y=para(text,y-14)-8
c.showPage();c.save()

# Word değerlendirmesinde gerçek başlık stilleri, tekrarlanan tablo başlıkları ve sabit sayfa bütçesi kullanır.
doc=Document();section=doc.sections[0];section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.9);section.left_margin=Cm(2.1);section.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(9);normal.paragraph_format.line_spacing=1.2
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(14);doc.styles['Heading 2'].paragraph_format.space_after=Pt(7);doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)
doc.styles['Title'].font.size=Pt(28);doc.styles['Heading 1'].font.size=Pt(21);doc.styles['Heading 2'].font.size=Pt(13)
doc.core_properties.author='Atlas Perakende';doc.core_properties.title='Atlas Perakende 2025 Ticari Değerlendirme';doc.core_properties.created=datetime(2025,12,31,12,tzinfo=timezone.utc);doc.core_properties.modified=datetime(2025,12,31,12,tzinfo=timezone.utc)
section.header.paragraphs[0].text='ATLAS PERAKENDE     /     TİCARİ DEĞERLENDİRME 2025'
section.header.paragraphs[0].runs[0].font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.text='Yönetim çalışma dosyası     ·     ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def dt(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            props=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E9EDF3' if ri==0 else ('F7F8FA' if ri%2 else 'FFFFFF'));props.append(shd)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            props.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
    doc.add_paragraph()

def dp(text):doc.add_paragraph(text)
def newpage(n,title,intro):doc.add_page_break();doc.add_paragraph(f'0{n}  /  2025 YÖNETİM DEĞERLENDİRMESİ',style='Subtitle');doc.add_heading(title,1);dp(intro)

doc.add_paragraph('Atlas Perakende 2025 ticari değerlendirme',style='Title');doc.add_paragraph('Yönetim ekibi için karar ve takip dosyası',style='Subtitle')
dp(f"2025 yılında {money(g['net'])} net satış üreten Atlas Perakende için öncelik, gelir kompozisyonunu katkı ve hizmet kalitesiyle birlikte yönetmektir. Bu değerlendirme dört bölge, üç kanal ve sekiz kategoriye ait aylık toplamları aynı çerçevede birleştirir.")
doc.add_heading('Kapsam ve kullanım',2);dp('Bu belge sayısal performans raporunun ticari yorumlama çerçevesini tanımlar. Satış kayıtları, aylık bütçe ve teslimat dosyası veri dayanağıdır. Sayısal bulgular gözlem; aşağıdaki karar başlıkları ise yönetimin ek veriyle değerlendireceği önerilerdir.')
dt(['Gösterge','2025'],[['Net satış',money(g['net'])],['Sipariş',num(g['orders'],0)],['Brüt marj','%'+num(g['margin'])],['Zamanında teslimat','%'+num(g['onTime'])]])
doc.add_heading('Okuma sırası',2);dp('Önce aylık bütçe ve gerçekleşme tablosunu, sonra kanal ekonomisini inceleyin. Kategori ve bölge kırılımlarıyla gelir bileşimini okuyun. Teslimat oranı ve mali kapsam sınırlamaları, kararın uygulanabilirliğini değerlendirirken birlikte kullanılmalıdır.')
newpage(2,'Kanal ekonomisi','Kanal performansının değerlendirilmesinde aynı döneme ait net satış, marj ve sipariş hacmi birlikte ele alınır. Kanal cirosu müşteri değeri veya kârlılık sıralaması yerine geçmez.')
for name in g['channel']:
    rs=[r for r in sales if r['channel']==name];n=sum(D(r['gross_sales_try'])-D(r['discount_try'])-D(r['refund_try']) for r in rs);p=n-sum(D(r['cost_of_goods_try']) for r in rs);orders=sum(int(r['orders']) for r in rs);con=p-sum(D(r['marketing_try'])+D(r['fulfillment_try']) for r in rs)
    doc.add_heading(name,2);dp(f"Net satış {money(n)}, brüt marj %{num(p/n*100)} ve sipariş başına katkı {money(con/orders)} olarak hesaplanmıştır. Kanalın büyümesi değerlendirilirken iade ve indirim tutarlarıyla kapasite ihtiyacı aynı dönem içinde incelenmelidir.")
doc.add_heading('Karar için ek kanıt',2);dp('Müşteri edinme maliyeti, tekrar satın alma, kanal komisyonu ve müşteri yaşam boyu değeri bu dosyada ayrı ölçülmemiştir. Bu nedenle kanallar arasında bütçe aktarımı otomatik sonuç olarak sunulmaz. Birim katkı genel gider ve finansman maliyetlerini içermez.')
newpage(3,'Portföy ve bölgesel odak','Gelir kırılımları birbirini tamamlar. Aynı satış bir kez kategori, bir kez bölge toplamına girer; farklı boyutların toplamlarını toplamak çift sayım üretir.')
dt(['Kategori','Yıllık net satış'],[[k,money(v)] for k,v in sorted(g['category'].items(),key=lambda x:-x[1])])
doc.add_heading('Bölgesel değerlendirme',2);dp('Marmara, Ege, İç Anadolu ve Akdeniz için satış ve teslimat kayıtları aynı ay ve kanal anahtarlarıyla ilişkilendirilir. Bölge büyüklüğü tek başına verimlilik göstergesi değildir; mağaza sayısı, alan, stok ve hizmet kapasitesiyle normalize edilmeden bölge performansı sıralanmaz.')
newpage(4,'Riskler ve kontrol noktaları','Aşağıdaki kontroller, yönetim değerlendirmesinin sayıların kapsamını koruması için uygulanır. Bunlar ticari riskin gerçekleştiğine dair bulgu değildir.')
for title,text in [('Kapsam tutarlılığı','Satış dosyası aylık kategori ve segment toplamıdır. Siparişler bu kırılımlarda ayrık kabul edilir. Tekil müşteri veya işlem sayısı olarak raporlanmaz. Tüm tutarlar TRY ve KDV hariçtir.'),('Dönem ve bütçe','Bütçe aynı ayın net satış hedefidir. Gerçekleşme oranında aylık oranların ortalaması kullanılmaz. Dönem filtresi satış ve bütçe kaynaklarına birlikte uygulanır.'),('Hizmet kalitesi','Zamanında ve gecikmeli siparişler toplamı sevk edilen siparişe eşittir. Satış ve operasyon siparişleri aylık bölge ve kanal anahtarında uzlaştırılır.'),('Yorum sınırı','Kampanya, fiyatlama veya kanal değişiminin nedensel etkisi bu tablolardan kanıtlanamaz. Karar önerisi gözlemden ayrı sunulur; dayanakları eksik karar için ek kayıt istenir.')]:doc.add_heading(title,2);dp(text)
newpage(5,'Karar planı ve takip','Önerilen iş sırası, önce sayısal kapsamı doğrulamayı, ardından karar için eksik bağlamı tamamlamayı hedefler. Bu iş sırası bir taahhüt değil, çalışma planı önerisidir.')
dt(['Çalışma','Sorumlu rol','Çıktı'],[['Kanal katkı incelemesi','Ticari finans','Birim katkı ve komisyon mutabakatı'],['Teslimat nedenleri','Operasyon','Bölge ve taşıyıcı neden kodları'],['Portföy kapasitesi','Kategori yönetimi','Stok ve tedarik süresi dökümü'],['Yatırım kararı','Yönetim ekibi','Varsayımları açık karar notu']])
doc.add_heading('İlk değerlendirme',2);dp('Mali kapsamı doğrulayın ve kanal bazında gider eşleştirmesini gözden geçirin. Toplamı açıklanamayan alanlarda gelir karşılaştırmasını durdurmak yerine alanı eksik olarak işaretleyin; doğrulanmış göstergeleri koruyun.')
doc.add_heading('Takip değerlendirmesi',2);dp('İade nedenleri, taşıyıcı performansı ve stok göstergeleri tamamlandığında karar notunu güncelleyin. Önceki dönemle karşılaştırmada aynı ürün ve kanal kapsamı kullanıldığını doğrulayın. Yeni veri olmadan belirsizliği sayısal güven puanına çevirmeyin.')
doc.add_heading('Kaynak eşleştirmesi',2);dp('atlas-satis-2025.csv satış ve mali kırılımları; atlas-butce-2025.csv aylık hedefleri; atlas-operasyon-2025.csv hizmet oranlarını taşır. atlas-performans-2025.pdf bu kayıtların yıllık yönetim görünümüdür. Bu belge yorumlama ve kontrol çerçevesini tamamlar.')
buffer=io.BytesIO();doc.save(buffer)
# DOCX ZIP zamanlarını normalize ederek aynı içerikte hash değişimini önler.
with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as zin, zipfile.ZipFile(ROOT/DOCX,'w',zipfile.ZIP_DEFLATED) as zout:
    for name in sorted(zin.namelist()):
        info=zipfile.ZipInfo(name,(2025,12,31,12,0,0));info.compress_type=zipfile.ZIP_DEFLATED;zout.writestr(info,zin.read(name))
manifest=json.loads((ROOT/'manifest.json').read_text())
manifest['files']=[i for i in manifest['files'] if i['file'].endswith('.csv')]
for source,file,pages in [('atlas-review',PDF,7),('atlas-commercial',DOCX,5)]:
    manifest['files'].append(dict(sourceId=source,file=file,pages=pages,sha256=hashlib.sha256((ROOT/file).read_bytes()).hexdigest()))
(ROOT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'rows':g['rows'],'net':g['net'],'pdfPages':7,'docxPages':5},ensure_ascii=False))
